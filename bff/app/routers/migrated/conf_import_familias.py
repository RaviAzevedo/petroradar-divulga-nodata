"""Rota MIGRADA (estrangulador): POST /api/conformidade/import/familias.

Espelha 1:1 o handler do Next:
  src/app/api/conformidade/import/familias/route.ts

Importação em massa de Famílias CRC via planilha XLSX.

  POST /api/conformidade/import/familias        -> preview (dry-run): retorna diff
  POST /api/conformidade/import/familias?aplicar=1 -> commit: aplica o diff

Body (multipart/form-data):  - file: XLSX da planilha mestre.

────────────────────────────────────────────────────────────────────────────────
AUTH (paridade EXATA com ``autenticar()`` do route.ts — réplica MANUAL, NÃO usa
``Depends(require_conformidade)`` porque a ORDEM e o caminho do Bearer divergem):

  1) Se há header ``Authorization: Bearer <token>``:
       - Trata o token como IMPORT-TOKEN do CLI (HS256, kind "import_cli",
         assinado com AUTH_SECRET) — NÃO como token de sessão NextAuth.
         Réplica de ``verificarImportToken`` (import-tokens.ts):
           * 3 segmentos, HS256, sig válida.
           * payload.kind == "import_cli".
           * payload.exp >= agora (em segundos).
         Inválido/expirado -> 401 {"error":"Token de importação inválido ou expirado"}.
       - Se válido: checa ``podeImportarConformidade(payload.role)`` (ADMIN/LIDER/
         SGI). Falha -> 403 {"error":"Token sem permissão de importação"}.
       - userId = payload.sub, role = payload.role.

  2) Senão (sem Bearer): sessão NextAuth via ``get_current_user``.
       - Sem sessão -> 401 {"error":"Não autenticado"}.
       - ``podeImportarConformidade(role)`` falha -> 403 {"error":"Permissão negada"}.

  IMPORTANTE: o Next NÃO normaliza o role do import-token (usa ``payload.role`` cru),
  enquanto o caminho de sessão usa ``session.user.role``. Para o import-token usamos
  o role CRU do payload (sem ``normalize_role``) — paridade com o Next, que passa
  ``payload.role`` direto a ``podeImportarConformidade``.

ORDEM DAS CHECAGENS (idêntica ao Next):
  auth(401/403) -> parse form (400 se não-multipart) -> file obrigatório (422) ->
  whitelist extensão (422) -> tamanho >10MB (422) -> magic-bytes (422) ->
  parse XLSX (422) -> diff -> (preview 200 | commit 200).

SHAPE de erro (cópia literal do route.ts — shape PLANO ``{"error": "<msg>"}``):
  - Não-multipart:      400 {"error":"Esperado multipart/form-data"}
  - file ausente:       422 {"error":"Arquivo XLSX é obrigatório"}
  - extensão proibida:  422 {"error":"Tipo de arquivo não permitido. Use XLSX ou XLS."}
  - > 10MB:             422 {"error":"Planilha muito grande (máx 10MB)"}
  - magic-bytes:        422 {"error":<msg da validação>}
  - falha no parser:    422 {"error":<msg do erro>}

NEGÓCIO (réplica de lib/conformidade/import-familias.ts):
  parsePlanilhaFamilias -> gerarDiffFamilias -> aplicarDiffFamilias.
  Bugs de PRODUÇÃO replicados fielmente:
    - gerarDiffFamilias: a query de existentes usa ``WHERE deletado_em IS NULL``
      e IGNORA ``codigosNorm`` (busca TODAS as famílias não-deletadas, não só as
      da planilha). Replicado tal-e-qual (SELECT sem filtro por código).
    - parser usa o XLSX com openpyxl (Next usa ExcelJS); a paridade é de DADOS
      (mesmas linhas/colunas/normalizações), não de motor.

DB:
  - aplicarDiffFamilias: INSERT em ``import_jobs`` (RETURNING id), loop de
    INSERT/UPDATE em ``familias`` (cada um com try/except que conta erro e
    continua — réplica do console.error + erros++), depois UPDATE final do job.
  - Tudo num único ``await db.commit()`` ao final (o Drizzle neon-http é
    autocommit; aqui commitamos explicitamente o conjunto).

SERIALIZAÇÃO (paridade JSON.parse(bff) == JSON.parse(next)):
  - uuid (job.id, familia.id) -> str(...).
  - validadeFim no preview -> ``.toISOString()`` do JS (UTC, ms, 'Z') via
    :func:`_js_to_iso_string` (a planilha sempre cria datas à meia-noite UTC).
  - jsonb (criterios/tags/payload) -> dict/list -> CAST(:p AS jsonb).
"""

from __future__ import annotations

import hashlib
import hmac
import json as _json
import re
import unicodedata
from datetime import datetime, timedelta, timezone
from typing import Any

import jwt
from fastapi import APIRouter, Depends, Request
from fastapi.responses import JSONResponse
from sqlalchemy import text

from app.core.config import settings
from app.core.security import (
    CurrentUser,
    get_current_user,
    is_conformidade,
)
from app.db.session import get_db

router = APIRouter()

# Extensões aceitas (EXTENSOES_PLANILHA do route.ts).
_EXTENSOES_PLANILHA = (".xlsx", ".xls")

# ── Colunas / mapas (réplica do import-familias.ts) ───────────────────────────
_COLUNAS_OBRIGATORIAS = ("codigoCrc", "nome")

_MAPA_CRITERIOS: list[tuple[str, str]] = [
    ("crit_juridico", "JURIDICO"),
    ("crit_economico", "ECONOMICO"),
    ("crit_tecnico", "TECNICO"),
    ("crit_tecnico_sms", "TECNICO_SMS"),
    ("crit_lfo", "LFO"),
    ("crit_revendedor", "REVENDEDOR"),
    ("crit_integridade", "INTEGRIDADE"),
]


# ════════════════════════════════════════════════════════════════════════════
# AUTH — réplica MANUAL de autenticar() (route.ts)
# ════════════════════════════════════════════════════════════════════════════


def _pode_importar_conformidade(role: str | None) -> bool:
    """Réplica de ``podeImportarConformidade`` -> ``isConformidade`` (permissions.ts).

    {ADMIN, LIDER, SGI}. ``is_conformidade`` aceita None com segurança (role inválido
    não está no conjunto).
    """
    return is_conformidade(role or "")


def _verificar_import_token(token: str) -> dict | None:
    """Réplica de ``verificarImportToken`` (import-tokens.ts).

    JWT HS256 assinado com AUTH_SECRET, payload com ``kind == "import_cli"`` e
    ``exp`` (segundos) ainda no futuro. Retorna o payload (dict) se válido, senão
    None. NÃO usa o caminho JWE de sessão — este é o token efêmero do CLI.
    """
    partes = token.split(".")
    if len(partes) != 3:
        return None

    secret = settings.jwt_secret
    if not secret:
        return None

    header_b64, payload_b64, sig_b64 = partes

    try:
        # Verificação HMAC-SHA256 manual (espelha crypto.subtle.verify do Next).
        esperado = hmac.new(
            secret.encode("utf-8"),
            f"{header_b64}.{payload_b64}".encode("utf-8"),
            hashlib.sha256,
        ).digest()
        recebido = _b64url_decode(sig_b64)
        if not hmac.compare_digest(esperado, recebido):
            return None

        payload = _json.loads(_b64url_decode(payload_b64).decode("utf-8"))
        if payload.get("kind") != "import_cli":
            return None
        exp = payload.get("exp")
        # payload.exp < now -> None (igual ao Next; comparação em segundos inteiros).
        agora = int(datetime.now(tz=timezone.utc).timestamp())
        if not isinstance(exp, (int, float)) or float(exp) < agora:
            return None
        return payload
    except Exception:  # noqa: BLE001 — qualquer falha de decode -> None (igual Next)
        return None


def _b64url_decode(s: str) -> bytes:
    """Decodifica base64url (sem padding), tolerante como o atob do Next."""
    import base64

    s = s.replace("-", "+").replace("_", "/")
    s += "=" * (-len(s) % 4)
    return base64.b64decode(s)


class _RequestSemAuthHeader:
    """Wrapper de ``Request`` que ESCONDE o header Authorization.

    Espelha ``auth()`` do NextAuth v5, que resolve a sessão SOMENTE pelo cookie e
    IGNORA totalmente o header ``Authorization``. ``get_current_user`` ->
    ``_extrair_token`` lê o header ANTES do cookie (e de forma case-insensitive,
    ``bearer ...``), então no caminho de sessão precisamos garantir que ele jamais
    consuma o header — caso contrário um header com casing não-canônico
    (``bearer <token-de-sessao>``) seria autenticado aqui, enquanto no Next cairia
    em ``auth()`` e só consultaria o cookie (401 "Não autenticado").

    Delegamos tudo ao ``Request`` real, mas devolvemos ``headers`` SEM a chave
    ``authorization``/``Authorization``, preservando os cookies. Assim
    ``_extrair_token`` cai direto no fallback de cookie, exatamente como ``auth()``.
    """

    def __init__(self, request: Request) -> None:
        self._request = request
        from starlette.datastructures import Headers

        filtrados = [
            (k, v)
            for k, v in request.headers.raw
            if k.lower() != b"authorization"
        ]
        self._headers = Headers(raw=filtrados)

    @property
    def headers(self):  # noqa: D401 — headers sem Authorization
        return self._headers

    def __getattr__(self, nome: str):
        # cookies, url, query_params, etc. vêm do Request real.
        return getattr(self._request, nome)


async def _autenticar(request: Request, db) -> dict:
    """Réplica de ``autenticar(request)`` do route.ts.

    Retorna ``{"userId": ..., "role": ...}`` em sucesso, ou
    ``{"erro": <msg>, "status": <code>}`` em falha (shape igual ao Next).
    """
    auth_header = (
        request.headers.get("authorization")
        or request.headers.get("Authorization")
        or ""
    )
    # Réplica EXATA do Next: o ramo de import-token só é tomado quando o esquema é
    # ``Bearer `` CASE-SENSITIVE (``authHeader.startsWith("Bearer ")``). Qualquer
    # outro casing (``bearer ``) NÃO entra aqui.
    if auth_header.startswith("Bearer "):
        token = auth_header[7:]
        payload = _verificar_import_token(token)
        if not payload:
            return {
                "erro": "Token de importação inválido ou expirado",
                "status": 401,
            }
        if not _pode_importar_conformidade(payload.get("role")):
            return {"erro": "Token sem permissão de importação", "status": 403}
        return {"userId": payload.get("sub"), "role": payload.get("role")}

    # Caminho de sessão NextAuth (== auth()): resolve a sessão SOMENTE pelo cookie.
    # Escondemos o header Authorization para que ``get_current_user`` NUNCA leia o
    # token do header (paridade com ``auth()``, que ignora o header). Assim, um
    # header ``bearer <token>`` com casing não-canônico cai em "sem sessão" -> 401
    # "Não autenticado", igual ao Next.
    try:
        user = await get_current_user(_RequestSemAuthHeader(request))
    except Exception:  # noqa: BLE001 — sessão ausente/ inválida -> "Não autenticado"
        return {"erro": "Não autenticado", "status": 401}

    if not _pode_importar_conformidade(user.role):
        return {"erro": "Permissão negada", "status": 403}
    return {"userId": user.id, "role": user.role}


# ════════════════════════════════════════════════════════════════════════════
# VALIDAÇÃO DE ARQUIVO — réplica de validarMagicBytes (storage/upload.ts)
# ════════════════════════════════════════════════════════════════════════════


def _is_pdf(b: bytes) -> bool:
    return len(b) >= 4 and b[0] == 0x25 and b[1] == 0x50 and b[2] == 0x44 and b[3] == 0x46


def _is_zip(b: bytes) -> bool:
    return len(b) >= 4 and b[0] == 0x50 and b[1] == 0x4B and b[2] == 0x03 and b[3] == 0x04


def _is_ole2(b: bytes) -> bool:
    return len(b) >= 4 and b[0] == 0xD0 and b[1] == 0xCF and b[2] == 0x11 and b[3] == 0xE0


def _validar_magic_bytes(conteudo: bytes) -> tuple[bool, str | None]:
    """Réplica de ``validarMagicBytes``: aceita PDF / ZIP(XLSX/DOCX) / OLE2(XLS/DOC)."""
    cabecalho = conteudo[:8]
    if _is_pdf(cabecalho) or _is_zip(cabecalho) or _is_ole2(cabecalho):
        return True, None
    return (
        False,
        "O conteúdo do arquivo não corresponde ao tipo declarado. Apenas PDF, "
        "XLS, XLSX, DOC e DOCX são permitidos.",
    )


# ════════════════════════════════════════════════════════════════════════════
# PARSING DE PLANILHA — réplica de parsePlanilhaFamilias (import-familias.ts)
# ════════════════════════════════════════════════════════════════════════════


def _str_trim(v: Any) -> str:
    """``String(v ?? "").trim()`` do JS (None -> "")."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        # ExcelJS entregaria um Date; o String(Date) do JS é a repr longa. Aqui só
        # importa para colunas de texto, que nunca recebem datas em planilha bem-
        # formada; conservamos o ISO para não quebrar caso aconteça.
        return v.isoformat()
    return str(v).strip()


def _parse_segmento(v: Any) -> str | None:
    """Réplica de ``parseSegmento``."""
    s = re.sub(r"\s+", "_", _str_trim(v).upper())
    if s in ("YELLOW_PAGES", "OPERACIONAL", "CRITICA", "ESTRATEGICA"):
        return s
    if s in ("AMARELAS", "PAGINAS_AMARELAS"):
        return "YELLOW_PAGES"
    if s == "OPERACIONAIS":
        return "OPERACIONAL"
    if s in ("CRÍTICA", "CRITICAS"):
        return "CRITICA"
    if s in ("ESTRATÉGICA", "ESTRATEGICAS"):
        return "ESTRATEGICA"
    return None


def _parse_status_fluxo(v: Any) -> str | None:
    """Réplica de ``parseStatusFluxo``."""
    s = _str_trim(v).upper()
    if s in ("D", "Q", "A", "X"):
        return s
    return None


def _parse_status_criterio(v: Any) -> str | None:
    """Réplica de ``parseStatusCriterio``."""
    s = _str_trim(v).upper()
    if s in ("APROVADO", "QUALIFICADO", "EM_AVALIACAO", "REJEITADO", "CANCELADO"):
        return s
    if s == "A":
        return "APROVADO"
    if s == "Q":
        return "QUALIFICADO"
    if s == "D":
        return "EM_AVALIACAO"
    if s == "R":
        return "REJEITADO"
    if s == "X":
        return "CANCELADO"
    return None


def _parse_date(v: Any) -> datetime | None:
    """Réplica de ``parseDate``: dd/mm/yyyy, yyyy-mm-dd, ou parse genérico.

    Retorna um ``datetime`` AWARE em UTC à meia-noite (espelha
    ``new Date("YYYY-MM-DDT00:00:00Z")``). Falha/None -> None.
    """
    if not v:
        return None
    if isinstance(v, datetime):
        return v
    # openpyxl pode entregar date (sem hora) para células de data.
    from datetime import date as _date

    if isinstance(v, _date):
        return datetime(v.year, v.month, v.day, tzinfo=timezone.utc)

    s = str(v).strip()
    if not s:
        return None
    m = re.match(r"^(\d{2})/(\d{2})/(\d{4})$", s)
    if m:
        return datetime(
            int(m.group(3)), int(m.group(2)), int(m.group(1)), tzinfo=timezone.utc
        )
    m2 = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", s)
    if m2:
        return datetime(
            int(m2.group(1)), int(m2.group(2)), int(m2.group(3)), tzinfo=timezone.utc
        )
    # Parse genérico (best-effort; ISO). new Date(s) inválido -> None.
    try:
        d = datetime.fromisoformat(s.replace("Z", "+00:00"))
        if d.tzinfo is None:
            d = d.replace(tzinfo=timezone.utc)
        return d
    except (ValueError, TypeError):
        return None


def _parse_tags(v: Any) -> list[str] | None:
    """Réplica de ``parseTags``: split por , ; | -> trim -> filter truthy."""
    if not v:
        return None
    s = str(v).strip()
    if not s:
        return None
    partes = [t.strip() for t in re.split(r"[,;|]", s)]
    return [t for t in partes if t]


def _parse_planilha_familias(conteudo: bytes) -> dict:
    """Réplica de ``parsePlanilhaFamilias``. Levanta ``ValueError`` em erro estrutural.

    Retorna ``{"linhas": [...], "erros": [...]}``. Cada linha é um dict com as
    mesmas chaves de ``LinhaFamiliaImport``.
    """
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    worksheets = wb.worksheets
    if not worksheets:
        raise ValueError("Planilha vazia")
    ws = worksheets[0]

    # Header na linha 1: mapeia rótulo -> índice de coluna (1-based).
    headers: dict[str, int] = {}
    primeira_linha = next(ws.iter_rows(min_row=1, max_row=1), ())
    for cell in primeira_linha:
        v = _str_trim(cell.value)
        if v:
            headers[v] = cell.column  # 1-based

    for obrig in _COLUNAS_OBRIGATORIAS:
        if not headers.get(obrig):
            raise ValueError(
                f'Coluna obrigatória "{obrig}" não encontrada no cabeçalho'
            )

    def get(row: tuple, name: str) -> Any:
        col = headers.get(name)
        if not col:
            return None
        idx = col - 1
        if idx < 0 or idx >= len(row):
            return None
        return row[idx].value

    linhas: list[dict] = []
    erros: list[dict] = []

    # Linhas de dados a partir da 2.
    i = 1
    for row in ws.iter_rows(min_row=2):
        i += 1
        codigo_crc_raw = _str_trim(get(row, "codigoCrc"))
        if not codigo_crc_raw:
            continue  # linha vazia

        nome = _str_trim(get(row, "nome"))
        if not nome:
            erros.append(
                {
                    "linhaPlanilha": i,
                    "codigoCrc": codigo_crc_raw,
                    "mensagem": "Nome obrigatório",
                }
            )
            continue

        criterios: dict[str, dict] = {}
        for coluna, key in _MAPA_CRITERIOS:
            status = _parse_status_criterio(get(row, coluna))
            if status:
                criterios[key] = {"status": status, "validadeFim": None}

        descricao = _str_trim(get(row, "descricao")) or None
        observacoes = _str_trim(get(row, "observacoes")) or None

        linhas.append(
            {
                "codigoCrc": codigo_crc_raw,
                "nome": nome,
                "descricao": descricao,
                "segmentoPetrobras": _parse_segmento(get(row, "segmentoPetrobras")),
                "statusFluxoCrc": _parse_status_fluxo(get(row, "statusFluxoCrc")),
                "validadeFim": _parse_date(get(row, "validadeFim")),
                "criterios": criterios,
                "tags": _parse_tags(get(row, "tags")),
                "observacoes": observacoes,
                "linhaPlanilha": i,
            }
        )

    wb.close()
    return {"linhas": linhas, "erros": erros}


# ── normalizarCodigoCrc (validations.ts) ──────────────────────────────────────


def _normalizar_codigo_crc(codigo: str) -> str:
    """Réplica de ``normalizarCodigoCrc``: NFD, remove diacríticos, só [A-Z0-9]."""
    nfd = unicodedata.normalize("NFD", codigo)
    sem_diacriticos = "".join(c for c in nfd if not unicodedata.combining(c))
    return re.sub(r"[^a-zA-Z0-9]", "", sem_diacriticos).upper()


# ════════════════════════════════════════════════════════════════════════════
# DIFF — réplica de gerarDiffFamilias (import-familias.ts)
# ════════════════════════════════════════════════════════════════════════════

# BUG DE PRODUÇÃO replicado: a query NÃO filtra por codigosNorm — só por
# deletado_em IS NULL (busca TODAS as famílias não-deletadas).
_SELECT_EXISTENTES_SQL = text(
    """
    SELECT
      id,
      codigo_normalizado,
      nome,
      descricao,
      segmento_petrobras,
      status_fluxo_crc,
      validade_fim,
      observacoes
    FROM familias
    WHERE deletado_em IS NULL
    """
)


def _iso_date_slice(dt: datetime | None) -> str:
    """``date.toISOString().slice(0, 10)`` do JS (UTC, 'YYYY-MM-DD'); None -> ""."""
    if dt is None:
        return ""
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    dt = dt.astimezone(timezone.utc)
    return f"{dt.year:04d}-{dt.month:02d}-{dt.day:02d}"


async def _gerar_diff_familias(db, linhas: list[dict], erros_parse: list[dict]) -> dict:
    """Réplica de ``gerarDiffFamilias``.

    Busca existentes (mesmo bug de produção: ignora os códigos), monta o mapa por
    código_normalizado e classifica cada linha em novas/atualizadas/inalteradas.
    """
    codigos_norm = [_normalizar_codigo_crc(l["codigoCrc"]) for l in linhas]

    existentes: list = []
    if len(codigos_norm) > 0:
        existentes = (
            await db.execute(_SELECT_EXISTENTES_SQL)
        ).mappings().all()

    mapa_exist = {e["codigo_normalizado"]: e for e in existentes}

    diff = {
        "novas": [],
        "atualizadas": [],
        "inalteradas": 0,
        "erros": list(erros_parse),
    }

    for linha in linhas:
        norm = _normalizar_codigo_crc(linha["codigoCrc"])
        exist = mapa_exist.get(norm)
        if not exist:
            diff["novas"].append(linha)
            continue

        mudancas: list[str] = []
        if (exist["nome"] or "") != linha["nome"]:
            mudancas.append("nome")
        if (exist["descricao"] or "") != (linha["descricao"] or ""):
            mudancas.append("descricao")
        if (exist["segmento_petrobras"] or "") != (linha["segmentoPetrobras"] or ""):
            mudancas.append("segmentoPetrobras")
        if (exist["status_fluxo_crc"] or "") != (linha["statusFluxoCrc"] or ""):
            mudancas.append("statusFluxoCrc")
        validade_atual_iso = _iso_date_slice(exist["validade_fim"])
        validade_nova_iso = _iso_date_slice(linha["validadeFim"])
        if validade_atual_iso != validade_nova_iso:
            mudancas.append("validadeFim")
        if (exist["observacoes"] or "") != (linha["observacoes"] or ""):
            mudancas.append("observacoes")

        if len(mudancas) == 0:
            diff["inalteradas"] += 1
        else:
            diff["atualizadas"].append(
                {"linha": linha, "existenteId": str(exist["id"]), "mudancas": mudancas}
            )

    return diff


# ════════════════════════════════════════════════════════════════════════════
# APPLY — réplica de aplicarDiffFamilias (import-familias.ts)
# ════════════════════════════════════════════════════════════════════════════

_INSERT_JOB_SQL = text(
    """
    INSERT INTO import_jobs (fonte, iniciado_por, status, total_arquivos, observacao)
    VALUES (:fonte, :iniciado_por, :status, :total_arquivos, :observacao)
    RETURNING id
    """
)

_INSERT_FAMILIA_SQL = text(
    """
    INSERT INTO familias (
      codigo_crc, codigo_normalizado, nome, descricao, nivel, path,
      cadastro_corporativo, segmento_petrobras, status, status_fluxo_crc,
      validade_fim, proxima_renovacao_em, criterios, tags, observacoes,
      criado_por, atualizado_por
    ) VALUES (
      :codigo_crc, :codigo_normalizado, :nome, :descricao, :nivel, :path,
      :cadastro_corporativo, :segmento_petrobras, :status, :status_fluxo_crc,
      :validade_fim, :proxima_renovacao_em, CAST(:criterios AS jsonb),
      CAST(:tags AS jsonb), :observacoes, :criado_por, :atualizado_por
    )
    """
)

_UPDATE_FAMILIA_SQL = text(
    """
    UPDATE familias SET
      nome = :nome,
      descricao = :descricao,
      segmento_petrobras = :segmento_petrobras,
      status_fluxo_crc = :status_fluxo_crc,
      validade_fim = :validade_fim,
      proxima_renovacao_em = :proxima_renovacao_em,
      observacoes = :observacoes,
      atualizado_por = :atualizado_por,
      atualizado_em = :atualizado_em
    WHERE id = :id
    """
)

_UPDATE_JOB_SQL = text(
    """
    UPDATE import_jobs SET
      status = :status,
      processados = :processados,
      novos = :novos,
      atualizados = :atualizados,
      erros = :erros,
      finalizado_em = :finalizado_em,
      payload = CAST(:payload AS jsonb)
    WHERE id = :id
    """
)


async def _aplicar_diff_familias(db, diff: dict, user_id: str) -> dict:
    """Réplica de ``aplicarDiffFamilias``.

    Insere o job, processa INSERTs e UPDATEs (cada um isolado: erro conta e segue),
    atualiza o job. Commit único ao final (o app usa neon-http autocommit).
    """
    job_row = (
        await db.execute(
            _INSERT_JOB_SQL,
            {
                "fonte": "PLANILHA",
                "iniciado_por": user_id,
                "status": "EM_ANDAMENTO",
                "total_arquivos": len(diff["novas"]) + len(diff["atualizadas"]),
                "observacao": "Importação de famílias CRC via planilha",
            },
        )
    ).mappings().first()
    job_id = str(job_row["id"])

    novos = 0
    atualizados = 0
    erros = len(diff["erros"])

    # ── INSERTs ──
    for l in diff["novas"]:
        try:
            validade_fim = l["validadeFim"]
            proxima_renovacao_em = (
                validade_fim - timedelta(days=60) if validade_fim else None
            )
            criterios_val = (
                _json.dumps(l["criterios"]) if len(l["criterios"]) > 0 else None
            )
            tags_val = _json.dumps(l["tags"]) if l["tags"] is not None else None

            # Subtransação (SAVEPOINT) para isolar a falha sem abortar o lote,
            # espelhando o try/catch por-linha do Drizzle (cada insert autônomo).
            async with db.begin_nested():
                await db.execute(
                    _INSERT_FAMILIA_SQL,
                    {
                        "codigo_crc": l["codigoCrc"],
                        "codigo_normalizado": _normalizar_codigo_crc(l["codigoCrc"]),
                        "nome": l["nome"],
                        "descricao": l["descricao"],
                        "nivel": "FAMILIA",
                        "path": l["codigoCrc"],
                        "cadastro_corporativo": True,
                        "segmento_petrobras": l["segmentoPetrobras"],
                        "status": "ATIVA",
                        "status_fluxo_crc": l["statusFluxoCrc"],
                        "validade_fim": validade_fim,
                        "proxima_renovacao_em": proxima_renovacao_em,
                        "criterios": criterios_val,
                        "tags": tags_val,
                        "observacoes": l["observacoes"],
                        "criado_por": user_id,
                        "atualizado_por": user_id,
                    },
                )
            novos += 1
        except Exception:  # noqa: BLE001 — réplica do catch{ erros++; console.error }
            erros += 1

    # ── UPDATEs ──
    for u in diff["atualizadas"]:
        try:
            linha = u["linha"]
            validade_fim = linha["validadeFim"]
            proxima_renovacao_em = (
                validade_fim - timedelta(days=60) if validade_fim else None
            )
            async with db.begin_nested():
                await db.execute(
                    _UPDATE_FAMILIA_SQL,
                    {
                        "id": u["existenteId"],
                        "nome": linha["nome"],
                        "descricao": linha["descricao"],
                        "segmento_petrobras": linha["segmentoPetrobras"],
                        "status_fluxo_crc": linha["statusFluxoCrc"],
                        "validade_fim": validade_fim,
                        "proxima_renovacao_em": proxima_renovacao_em,
                        "observacoes": linha["observacoes"],
                        "atualizado_por": user_id,
                        "atualizado_em": datetime.now(tz=timezone.utc),
                    },
                )
            atualizados += 1
        except Exception:  # noqa: BLE001 — réplica do catch{ erros++; console.error }
            erros += 1

    # ── UPDATE final do job ──
    status_final = (
        "FALHOU" if (erros > 0 and novos + atualizados == 0) else "CONCLUIDO"
    )
    payload = {
        "novosCodigos": [l["codigoCrc"] for l in diff["novas"]],
        "atualizadosCodigos": [u["linha"]["codigoCrc"] for u in diff["atualizadas"]],
        "erros": diff["erros"],
    }
    await db.execute(
        _UPDATE_JOB_SQL,
        {
            "id": job_id,
            "status": status_final,
            "processados": novos + atualizados,
            "novos": novos,
            "atualizados": atualizados,
            "erros": erros,
            "finalizado_em": datetime.now(tz=timezone.utc),
            "payload": _json.dumps(payload),
        },
    )

    await db.commit()

    return {
        "jobId": job_id,
        "novos": novos,
        "atualizados": atualizados,
        "erros": erros,
    }


# ── Serialização do preview ───────────────────────────────────────────────────


def _js_to_iso_string(dt: datetime | None) -> str | None:
    """Réplica de ``new Date(dt).toISOString()`` (UTC, ms, sufixo 'Z')."""
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    dt = dt.astimezone(timezone.utc)
    ms = dt.microsecond // 1000
    return (
        f"{dt.year:04d}-{dt.month:02d}-{dt.day:02d}"
        f"T{dt.hour:02d}:{dt.minute:02d}:{dt.second:02d}.{ms:03d}Z"
    )


# ════════════════════════════════════════════════════════════════════════════
# HANDLER POST
# ════════════════════════════════════════════════════════════════════════════


@router.post("/api/conformidade/import/familias")
async def import_familias(
    request: Request,
    db=Depends(get_db),
) -> JSONResponse:
    """POST: preview (dry-run) ou commit (``?aplicar=1``) da importação de famílias.

    Réplica fiel do handler POST do route.ts: mesma ordem de checagens, mesmos
    status/shapes de erro, mesma lógica de diff/apply e mesmo formato de resposta.
    """
    # 1) AUTH — réplica de autenticar(): 401/403 com shape {"error": <msg>}.
    a = await _autenticar(request, db)
    if "erro" in a:
        return JSONResponse(status_code=a["status"], content={"error": a["erro"]})

    aplicar = request.query_params.get("aplicar") == "1"

    # 2) Parse multipart -> 400 se não for multipart/form-data.
    try:
        form = await request.form()
    except Exception:  # noqa: BLE001
        return JSONResponse(
            status_code=400, content={"error": "Esperado multipart/form-data"}
        )

    file = form.get("file")
    # Em FastAPI/Starlette, um campo de arquivo é UploadFile; texto é str. O Next
    # exige `file instanceof File` -> qualquer não-arquivo (ausente/texto) cai no 422.
    from starlette.datastructures import UploadFile as _UploadFile

    if not isinstance(file, _UploadFile):
        return JSONResponse(
            status_code=422, content={"error": "Arquivo XLSX é obrigatório"}
        )

    # 3) Whitelist de extensão (.xlsx / .xls).
    nome_arquivo = file.filename or ""
    partes_nome = nome_arquivo.split(".")
    ext = "." + (partes_nome[-1].lower() if len(partes_nome) > 1 else "")
    if ext not in _EXTENSOES_PLANILHA:
        return JSONResponse(
            status_code=422,
            content={"error": "Tipo de arquivo não permitido. Use XLSX ou XLS."},
        )

    # Lê o conteúdo (necessário para tamanho, magic-bytes e parser).
    conteudo = await file.read()

    # 4) Tamanho > 10MB.
    if len(conteudo) > 10 * 1024 * 1024:
        return JSONResponse(
            status_code=422, content={"error": "Planilha muito grande (máx 10MB)"}
        )

    # 5) Magic-bytes.
    valido, erro_magic = _validar_magic_bytes(conteudo)
    if not valido:
        return JSONResponse(
            status_code=422,
            content={
                "error": erro_magic
                or "Conteúdo do arquivo não corresponde ao tipo declarado"
            },
        )

    # 6) Parse da planilha -> 422 com a mensagem do erro.
    try:
        parsed = _parse_planilha_familias(conteudo)
    except Exception as e:  # noqa: BLE001
        msg = str(e) if str(e) else "Falha ao processar planilha"
        return JSONResponse(status_code=422, content={"error": msg})

    # 7) Diff.
    diff = await _gerar_diff_familias(db, parsed["linhas"], parsed["erros"])

    # 8) Preview vs Commit.
    if not aplicar:
        return JSONResponse(
            content={
                "modo": "preview",
                "resumo": {
                    "novas": len(diff["novas"]),
                    "atualizadas": len(diff["atualizadas"]),
                    "inalteradas": diff["inalteradas"],
                    "erros": len(diff["erros"]),
                },
                "novas": [
                    {
                        "linhaPlanilha": l["linhaPlanilha"],
                        "codigoCrc": l["codigoCrc"],
                        "nome": l["nome"],
                        "segmento": l["segmentoPetrobras"],
                        "statusFluxo": l["statusFluxoCrc"],
                        "validadeFim": _js_to_iso_string(l["validadeFim"]),
                    }
                    for l in diff["novas"]
                ],
                "atualizadas": [
                    {
                        "linhaPlanilha": u["linha"]["linhaPlanilha"],
                        "codigoCrc": u["linha"]["codigoCrc"],
                        "nome": u["linha"]["nome"],
                        "mudancas": u["mudancas"],
                    }
                    for u in diff["atualizadas"]
                ],
                "erros": diff["erros"],
            }
        )

    # Commit.
    resultado = await _aplicar_diff_familias(db, diff, a["userId"])
    return JSONResponse(content={"modo": "commit", **resultado})
