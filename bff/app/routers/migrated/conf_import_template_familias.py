"""Rota migrada (strangler-fig): GET /api/conformidade/import/template/familias.

Equivalente 1:1 de
``src/app/api/conformidade/import/template/familias/route.ts`` (handler GET).

PARTICULARIDADE DESTA ROTA (NÃO é uma rota de leitura de banco):
  O handler do Next NÃO toca o banco. Ele monta, em memória, o **template XLSX**
  que o SGI deve preencher e devolve o BINÁRIO do arquivo como download
  (``Content-Type`` de planilha + ``Content-Disposition: attachment``). Logo:

    - Não há ``db`` (sem ``Depends(get_db)``), sem query, sem query params.
    - O corpo NÃO é JSON — é o arquivo .xlsx. A regra de paridade
      "JSON.parse(bff) == JSON.parse(next)" não se aplica ao corpo binário;
      a paridade aqui é estrutural: MESMAS planilhas, MESMO cabeçalho, MESMA
      linha de exemplo, MESMOS comentários, MESMAS larguras, MESMO painel
      congelado, e MESMOS headers/guard de auth.

GUARD (paridade exata com o Next):
  Next: ``auth()`` (sessão) + ``podeImportarConformidade(role)``.
  ``podeImportarConformidade`` → ``isConformidade`` → papéis {ADMIN, LIDER, SGI}.
  No BFF isso é exatamente :data:`app.core.security.require_conformidade`
  (ROLES_CONFORMIDADE = ADMIN, LIDER, SGI). Sessão ausente → 401 (via
  get_current_user); papel fora do conjunto → 403 (via require_role).

RESPOSTA (paridade exata com o Next):
  status 200; body = bytes do .xlsx; headers:
    Content-Type:        application/vnd.openxmlformats-officedocument.spreadsheetml.sheet
    Content-Disposition: attachment; filename="familias_template.xlsx"
    Cache-Control:       no-store

IMPLEMENTAÇÃO:
  O Next usa ExcelJS; aqui usamos ``openpyxl`` (já disponível no BFF) para
  reproduzir a MESMA estrutura/estilos do arquivo gerado. Cores em ARGB do
  ExcelJS (``FF1D1D1B`` etc.) mapeiam direto para o ``aRGB`` do openpyxl.
"""

from __future__ import annotations

import io

from fastapi import APIRouter, Depends, Request
from fastapi.responses import Response

from app.core.security import CurrentUser, require_conformidade

router = APIRouter()


# ── Colunas do template (réplica EXATA do array COLUNAS do route.ts) ───────────
# Cada tupla: (key/label, largura, obrigatorio, exemplo, comentario)
_COLUNAS: list[tuple[str, int, bool, str, str | None]] = [
    ("codigoCrc", 14, True, "G02.01.05", None),
    ("nome", 40, True, "Bombas centrífugas", None),
    ("descricao", 50, False, "Descrição livre", None),
    ("segmentoPetrobras", 18, False, "CRITICA",
     "YELLOW_PAGES | OPERACIONAL | CRITICA | ESTRATEGICA"),
    ("statusFluxoCrc", 14, False, "A",
     "D=Declaração | Q=Qualificado | A=Aprovado | X=Cancelado"),
    ("validadeFim", 14, False, "2027-03-15",
     "Data ISO YYYY-MM-DD ou DD/MM/YYYY"),
    ("crit_juridico", 12, False, "A",
     "A=Aprovado | Q=Qualificado | D=Em Avaliação | R=Rejeitado | X=Cancelado"),
    ("crit_economico", 12, False, "A", None),
    ("crit_tecnico", 12, False, "A", None),
    ("crit_tecnico_sms", 14, False, "A", None),
    ("crit_lfo", 10, False, "A", None),
    ("crit_revendedor", 14, False, "", None),
    ("crit_integridade", 14, False, "A", None),
    ("tags", 24, False, "petroquimica;onshore", None),
    ("observacoes", 40, False, "", None),
]

# ── Instruções (réplica EXATA do addRows da sheet "Instruções") ───────────────
_INSTRUCOES: list[tuple[str, str]] = [
    ("codigoCrc",
     "Código oficial Petrobras (ex.: G02.01.05). Obrigatório. Único. Aceita pontuação (será normalizada)."),
    ("nome",
     "Nome da família. Obrigatório. Ex.: 'Bombas centrífugas'."),
    ("descricao",
     "Descrição livre do escopo coberto. Opcional."),
    ("segmentoPetrobras",
     "Classificação Petrobras: YELLOW_PAGES, OPERACIONAL, CRITICA, ESTRATEGICA. Aceita também 'Páginas Amarelas', 'Crítica' etc."),
    ("statusFluxoCrc",
     "Status oficial Petrobras: D=Declaração de Interesse, Q=Tecnicamente Qualificado, A=Aprovado, X=Cancelado."),
    ("validadeFim",
     "Data de vencimento do CRC nesta família. Aceita YYYY-MM-DD ou DD/MM/YYYY. Sistema calcula 'proximaRenovacaoEm' automaticamente (60 dias antes)."),
    ("crit_*",
     "Status individual de cada critério Petrobras. Use letras A/Q/D/R/X ou nomes APROVADO/QUALIFICADO/EM_AVALIACAO/REJEITADO/CANCELADO. Critérios oficiais: Jurídico, Econômico, Técnico, Técnico-SMS, LFO (Linha de Fornecimento), Revendedor (apenas distribuidores), Integridade."),
    ("tags",
     "Lista separada por ; , ou |. Ex.: 'petroquimica;onshore;manutencao'."),
    ("observacoes",
     "Texto livre para anotações internas SGI."),
]

# Tipo MIME e nome de arquivo (idênticos ao route.ts).
_XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
_FILENAME = "familias_template.xlsx"


def _construir_workbook() -> bytes:
    """Monta o template XLSX (mesma estrutura do ExcelJS do Next) e devolve bytes.

    Réplica fiel do route.ts:
      - Sheet "Familias": cabeçalho estilizado (preto #1D1D1B, fonte branca bold),
        obrigatórios em laranja #F59C00, comentários nas células com 'comentario',
        linha de exemplo em itálico cinza #94A3B8, painel congelado abaixo da
        linha 1, larguras por coluna.
      - Sheet "Instruções": cabeçalho preto/branco bold, larguras 22/90, e as
        linhas Campo/Descrição.
    """
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.properties.creator = "Petro Radar — Conformidade"

    # ── Sheet principal: "Familias" ──
    ws = wb.active
    ws.title = "Familias"

    preto = PatternFill(fill_type="solid", fgColor="FF1D1D1B")
    laranja = PatternFill(fill_type="solid", fgColor="FFF59C00")
    fonte_header = Font(bold=True, color="FFFFFFFF", size=11)
    align_header = Alignment(vertical="center", horizontal="left")
    fonte_exemplo = Font(italic=True, color="FF94A3B8")

    # Cabeçalho (linha 1): rótulo, largura, estilo, obrigatório (laranja), nota.
    for idx, (key, largura, obrigatorio, _exemplo, comentario) in enumerate(
        _COLUNAS, start=1
    ):
        letra = get_column_letter(idx)
        ws.column_dimensions[letra].width = largura
        celula = ws.cell(row=1, column=idx, value=key)
        celula.font = fonte_header
        celula.alignment = align_header
        celula.fill = laranja if obrigatorio else preto
        if comentario:
            celula.comment = Comment(comentario, "Petro Radar")
    ws.row_dimensions[1].height = 22

    # Linha de exemplo (linha 2): valores 'exemplo', fonte itálica cinza.
    for idx, (_key, _largura, _obr, exemplo, _com) in enumerate(_COLUNAS, start=1):
        celula = ws.cell(row=2, column=idx, value=exemplo if exemplo else None)
        celula.font = fonte_exemplo

    # Painel congelado abaixo do cabeçalho (ySplit:1 do ExcelJS == freeze "A2").
    ws.freeze_panes = "A2"

    # ── Sheet de instruções: "Instruções" ──
    ws_inst = wb.create_sheet("Instruções")
    ws_inst.column_dimensions["A"].width = 22
    ws_inst.column_dimensions["B"].width = 90
    for col, titulo in ((1, "Campo"), (2, "Descrição")):
        c = ws_inst.cell(row=1, column=col, value=titulo)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = preto
    for linha, (campo, desc) in enumerate(_INSTRUCOES, start=2):
        ws_inst.cell(row=linha, column=1, value=campo)
        ws_inst.cell(row=linha, column=2, value=desc)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@router.get("/api/conformidade/import/template/familias")
async def baixar_template_familias(
    request: Request,
    user: CurrentUser = Depends(require_conformidade),
) -> Response:
    """GET /api/conformidade/import/template/familias — template XLSX de famílias.

    Guard paritário: ``require_conformidade`` (ADMIN/LIDER/SGI), espelhando
    ``auth()`` + ``podeImportarConformidade(role)`` do Next (401 sem sessão,
    403 para papel fora do conjunto). Devolve o BINÁRIO do .xlsx com os MESMOS
    headers do route.ts (Content-Type de planilha, attachment, no-store).
    """
    conteudo = _construir_workbook()
    return Response(
        content=conteudo,
        status_code=200,
        media_type=_XLSX_MIME,
        headers={
            "Content-Type": _XLSX_MIME,
            "Content-Disposition": f'attachment; filename="{_FILENAME}"',
            "Cache-Control": "no-store",
        },
    )
